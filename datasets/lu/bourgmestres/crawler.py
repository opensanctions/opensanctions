from datetime import datetime
from typing import Any

from openpyxl import load_workbook
from rigour.mime.types import XLSX

from zavod import Context
from zavod import helpers as h
from zavod.stateful.positions import categorise

DATASETS_QUERY_URL = "https://data.public.lu/api/1/datasets/?lang=en&q=bourgmestre"


def to_source_date(value: str | None) -> str | None:
    if value is None:
        return None
    return datetime.strptime(value, "%Y-%m-%d").strftime("%d/%m/%Y")


def crawl_record(
    context: Context,
    record: dict[str, Any],
) -> None:
    commune_code = record.pop("com_code")
    commune_label = record.pop("com_label")
    position_label = record.pop("man_label")
    first_name = record.pop("elu_first_name")
    last_name = record.pop("elu_last_name")
    gender = record.pop("elu_sex")
    # Render dates as the retired JSON resource did (dd/mm/YYYY): make_occupancy hashes the
    # raw date strings into the Occupancy ID, so ISO dates would re-key every occupancy.
    start_date = to_source_date(record.pop("coac_start_date"))
    end_date = to_source_date(record.pop("coac_end_date"))

    res = context.lookup("position", position_label)
    assert res is not None, f"Unknown position: {position_label!r}"
    position_name = res.value

    position = h.make_position(
        context,
        name=f"{position_name} of {commune_label}",
        country="lu",
        subnational_area=commune_label,
        topics=["gov.muni", "gov.executive"],
        lang="fra",
    )
    categorisation = categorise(context, position)
    if not categorisation.is_pep:
        return

    person = context.make("Person")
    person.id = context.make_id(first_name, last_name, commune_code)
    h.apply_name(person, first_name=first_name, last_name=last_name, lang="fra")
    person.add("gender", gender)
    # citizenship not required: https://guichet.public.lu/fr/citoyens/citoyennete/elections/elections-communales/candidat-elections-communales.html
    person.add("country", "lu")

    occupancy = h.make_occupancy(
        context,
        person,
        position,
        categorisation=categorisation,
        start_date=start_date,
        end_date=end_date or None,
    )
    if occupancy is not None:
        context.emit(position)
        context.emit(occupancy)
        context.emit(person)

    context.audit_data(record)


def crawl(context: Context) -> None:
    response = context.fetch_json(DATASETS_QUERY_URL, cache_days=1)
    assert isinstance(response, dict)

    # If there is more than one dataset for our search for "bourgmestre", a new election may have happened.
    total = response["total"]
    if total != 1:
        context.log.warning(
            "Expected exactly one dataset, found %d — a new election may have happened",
            total,
        )
        raise
    dataset = response["data"][0]

    # If the dataset page URL has changed, something changed, we want to bail
    if dataset["page"] != context.data_url:
        context.log.error(
            "Dataset page URL has changed — a new election may have happened: %s",
            dataset["page"],
        )
        raise

    # The publisher replaced the JSON resource with CSV + XLSX on 2026-10-14. The CSV is
    # Latin-1 and mangles names outside it (e.g. "Ćosić" -> "?osi?"), so read the XLSX.
    # The same file is sometimes re-uploaded as a second XLSX resource: take the newest.
    xlsx_resources = [r for r in dataset["resources"] if r["format"] == "xlsx"]
    assert len(xlsx_resources) > 0, "No XLSX resource found in dataset"
    xlsx_resource = max(xlsx_resources, key=lambda r: r["last_modified"])

    path = context.fetch_resource("source.xlsx", xlsx_resource["url"])
    context.export_resource(path, XLSX, title=context.SOURCE_TITLE)
    workbook = load_workbook(path, read_only=True)
    # The sheet name carries the upload date (e.g. "Schäfferot20261014").
    assert len(workbook.worksheets) == 1, workbook.sheetnames
    for record in h.parse_xlsx_sheet(context, workbook.worksheets[0]):
        crawl_record(context, record)
